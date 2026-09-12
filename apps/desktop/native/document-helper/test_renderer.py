"""Semantic regression tests; run with the pinned document-helper Python."""
import base64
import atexit
import importlib.util
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET

HERE = Path(__file__).resolve().parent
WORK = tempfile.TemporaryDirectory(prefix='zq-renderer-tests-')
atexit.register(WORK.cleanup)
WORK_PATH = Path(WORK.name)
shutil.copyfile(HERE / 'renderer.py', WORK_PATH / 'renderer.py')
font_spec = importlib.util.spec_from_file_location('prepare_fonts', HERE / 'prepare-fonts.py')
prepare_fonts = importlib.util.module_from_spec(font_spec)
font_spec.loader.exec_module(prepare_fonts)
FONT_SOURCE = HERE.parents[3] / 'node_modules' / '@fontsource' / 'noto-sans' / 'files'
(WORK_PATH / 'fonts').mkdir()
for subset in prepare_fonts.SUBSETS:
    name = 'noto-sans-' + subset + '-400-normal'
    (WORK_PATH / 'fonts' / (name + '.ttf')).write_bytes(prepare_fonts.convert((FONT_SOURCE / (name + '.woff')).read_bytes()))
spec = importlib.util.spec_from_file_location('renderer', WORK_PATH / 'renderer.py')
renderer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(renderer)


def request(format='docx', blocks=None, **typography):
    return {'version': 1, 'operation': 'render_document', 'document': {
        'format': format, 'title': 'Quarterly résumé', 'typography': typography,
        'blocks': blocks if blocks is not None else [
            {'type': 'heading', 'level': 2, 'text': 'Overview'},
            {'type': 'paragraph', 'text': 'Café Ελληνικά Привет\nSecond line'},
            {'type': 'code', 'text': 'first()\n  second()'},
            {'type': 'list', 'ordered': True, 'marker': '3.', 'text': 'Third item'},
            {'type': 'table', 'rows': [['Name', 'Value'], ['alpha', '=SUM(A1:A2)']]},
        ]}}


def render(value):
    result = renderer.handle_request(json.dumps(value).encode())
    if not result.get('ok'):
        raise AssertionError(result)
    return io.BytesIO(base64.b64decode(result['file']['data']))


class ValidationTests(unittest.TestCase):
    def reject(self, value, contains=None):
        result = renderer.handle_request(json.dumps(value).encode())
        self.assertEqual(set(result), {'version', 'ok', 'error'})
        self.assertEqual(result['version'], 1)
        self.assertIs(result['ok'], False)
        self.assertIsInstance(result['error'], str)
        if contains:
            self.assertIn(contains, result['error'])

    def test_rejects_unknown_keys_at_every_boundary(self):
        for where in ('request', 'document', 'typography', 'block'):
            value = request()
            target = {'request': value, 'document': value['document'],
                      'typography': value['document']['typography'],
                      'block': value['document']['blocks'][0]}[where]
            target['path'] = '/tmp/untrusted.py'
            with self.subTest(where=where):
                self.reject(value)

    def test_rejects_malformed_protocol_and_typography(self):
        for field, bad in [('version', True), ('version', 2), ('operation', 'exec')]:
            value = request(); value[field] = bad; self.reject(value)
        for typo in ({'bodySize': True}, {'bodySize': 25}, {'headingSize': 7},
                     {'fontFamily': '/tmp/font.ttf'}, {'bodySize': float('nan')}):
            value = request(); value['document']['typography'] = typo; self.reject(value)
        for bad in (None, [], {}, {'type': 'table', 'rows': []},
                    {'type': 'table', 'rows': [['a'], ['b', 'c']]},
                    {'type': 'heading', 'level': True, 'text': 'bad'},
                    {'type': 'paragraph', 'text': '\ud800'},
                    {'type': 'paragraph', 'text': '\x00bad'}):
            value = request(blocks=[bad]); self.reject(value)

    def test_rejects_source_request_block_row_and_column_limits(self):
        self.assertFalse(renderer.handle_request(b' ' * (1024 * 1024 + 1))['ok'])
        self.reject(request(blocks=[{'type': 'paragraph', 'text': 'é' * 51201}]), '100 KB')
        self.reject(request(blocks=[{'type': 'paragraph', 'text': 'a'}] * 4001), '4000')
        self.reject(request(blocks=[{'type': 'table', 'rows': [['a']] * 2001}]), '2000')
        self.reject(request(blocks=[{'type': 'table', 'rows': [['a'] * 31]}]), '30')

    def test_duplicate_keys_and_invalid_json_fail_closed(self):
        for raw in (b'{"version":1,"version":1}', b'NaN', b'null', b'[]', b'\xff', b'{'):
            self.assertFalse(renderer.handle_request(raw)['ok'])

    def test_safe_filename_and_exact_success_contract(self):
        value = request(); value['document']['title'] = '../a\\b: c'
        result = renderer.handle_request(json.dumps(value).encode())
        self.assertEqual(set(result), {'version', 'ok', 'file', 'warnings'})
        self.assertEqual(set(result['file']), {'name', 'mime', 'data'})
        self.assertEqual(result['file']['name'], '-a-b- c.docx')
        self.assertEqual(result['warnings'], [])

    def test_unicode_filename_obeys_transport_byte_limit(self):
        value = request(); value['document']['title'] = '文' * 160
        result = renderer.handle_request(json.dumps(value).encode())
        self.assertTrue(result['ok'])
        self.assertLessEqual(len(result['file']['name'].encode('utf-8')), 256)

    def test_cli_supports_supervisor_request_file_and_stdin(self):
        raw = json.dumps(request('xlsx')).encode()
        path = WORK_PATH / 'request.json'; path.write_bytes(raw)
        for args, stdin in (([str(path)], None), ([], raw)):
            completed = subprocess.run([sys.executable, '-I', '-B', str(WORK_PATH / 'renderer.py'), *args],
                                       input=stdin, capture_output=True, timeout=30, check=True)
            self.assertTrue(json.loads(completed.stdout)['ok'])
            self.assertEqual(completed.stderr, b'')

    def test_output_buffer_stops_oversized_documents(self):
        buffer = renderer.OutputBuffer()
        with self.assertRaisesRegex(renderer.DocumentError, '4 MB'):
            buffer.write(b'x' * (4 * 1024 * 1024 + 1))
        self.assertEqual(buffer.getvalue(), b'')

    def test_prepared_fonts_have_valid_sfnt_checksum(self):
        import struct
        for path in (WORK_PATH / 'fonts').glob('*.ttf'):
            data = path.read_bytes()
            self.assertEqual(sum(struct.unpack('>' + 'I' * (len(data) // 4), data)) & 0xFFFFFFFF, 0xB1B0AFBA)
        with self.assertRaises(ValueError):
            prepare_fonts.convert(b'not a font')


class DocumentTests(unittest.TestCase):
    def test_docx_export_writes_role_fonts_directly_without_template_theme_overrides(self):
        exported = render(request(titleFontFamily='Brush Script MT', headingFontFamily='Georgia',
                                  bodyFontFamily='Bradley Hand', titleSize=30, headingSize=22, bodySize=18))
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        q = lambda name: '{' + ns['w'] + '}' + name
        with zipfile.ZipFile(exported) as archive:
            document = ET.fromstring(archive.read('word/document.xml'))
            styles = ET.fromstring(archive.read('word/styles.xml'))
        for run in document.findall('.//w:r', ns):
            text = ''.join(node.text or '' for node in run.findall('w:t', ns))
            if not text:
                continue
            font, size = ('Brush Script MT', 60) if text == 'Quarterly résumé' else (
                ('Georgia', 40) if text == 'Overview' else
                ('Courier New', 36) if text.startswith('first()') else ('Bradley Hand', 36))
            fonts = run.find('w:rPr/w:rFonts', ns)
            self.assertIsNotNone(fonts, f'Missing explicit run font for {text}')
            for attribute in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                self.assertEqual(fonts.get(q(attribute)), font, text)
            self.assertEqual(run.find('w:rPr/w:sz', ns).get(q('val')), str(size), text)
            self.assertEqual(run.find('w:rPr/w:color', ns).get(q('val')), '202124', text)
        for element in styles.iter():
            self.assertFalse(any('theme' in key.lower() for key in element.attrib), element.attrib)
        defaults = styles.find('w:docDefaults/w:rPrDefault/w:rPr', ns)
        self.assertEqual(defaults.find('w:rFonts', ns).get(q('ascii')), 'Bradley Hand')
        self.assertEqual(defaults.find('w:sz', ns).get(q('val')), '36')

    def test_docx_retains_text_tables_line_breaks_and_role_styles(self):
        from docx import Document
        document = Document(render(request(titleFontFamily='Georgia', bodyFontFamily='Verdana',
                                           headingSize=22, bodySize=12)))
        self.assertEqual(document.paragraphs[0].text, 'Quarterly résumé')
        self.assertIn('Café Ελληνικά Привет\nSecond line', [p.text for p in document.paragraphs])
        self.assertIn('first()\n  second()', [p.text for p in document.paragraphs])
        self.assertEqual(document.tables[0].cell(1, 1).text, '=SUM(A1:A2)')
        self.assertEqual(document.styles['Title'].font.name, 'Georgia')
        self.assertEqual(document.styles['Normal'].font.name, 'Verdana')
        self.assertEqual(document.styles['Heading 2'].font.size.pt, 20)
        self.assertEqual(document.styles['Normal'].font.size.pt, 12)

    def test_xlsx_never_turns_literal_input_into_formulas(self):
        from openpyxl import load_workbook
        literals = ['=SUM(A1:A2)', '+cmd', '-1+1', '@SUM(A1)', '#N/A', '00123']
        sheet = load_workbook(render(request('xlsx', [{'type': 'table', 'rows': [literals]}],
                                            fontFamily='Georgia', bodySize=13))).active
        self.assertEqual([c.value for c in sheet[2]], literals)
        for cell in sheet[2]:
            self.assertEqual(cell.data_type, 's')
            self.assertEqual(cell.number_format, '@')
            self.assertEqual(cell.font.name, 'Georgia')
        self.assertEqual(sheet.freeze_panes, 'A2')

    def test_xlsx_does_not_read_machine_mime_configuration(self):
        foreign = WORK_PATH / 'forbidden.mime.types'; foreign.write_text('text/plain txt\n')
        script = '''
import builtins, json, mimetypes, runpy, sys
original_open = builtins.open
def guarded_open(path, *args, **kwargs):
    if str(path) == sys.argv[2]:
        raise PermissionError('Foreign machine MIME configuration must not be read')
    return original_open(path, *args, **kwargs)
builtins.open = guarded_open
mimetypes.knownfiles = [sys.argv[2]]
worker = runpy.run_path(sys.argv[1])
print(json.dumps(worker['handle_request'](sys.stdin.buffer.read())))
'''
        completed = subprocess.run([sys.executable, '-I', '-B', '-c', script,
                                    str(WORK_PATH / 'renderer.py'), str(foreign)],
                                   input=json.dumps(request('xlsx')).encode(), capture_output=True,
                                   timeout=30, check=True)
        self.assertTrue(json.loads(completed.stdout)['ok'], completed.stdout)
        self.assertEqual(completed.stderr, b'')

    def test_xlsx_line_rows_and_cell_limits(self):
        from openpyxl import load_workbook
        sheet = load_workbook(render(request('xlsx'))).active
        self.assertEqual(sheet['A4'].value, 'Second line')
        for blocks in ([{'type': 'paragraph', 'text': 'x' * 32768}],
                       [{'type': 'code', 'text': '\n'.join(['x'] * 2000)}]):
            self.assertFalse(renderer.handle_request(json.dumps(request('xlsx', blocks)).encode())['ok'])

    def test_pptx_has_editable_tables_text_and_role_fonts(self):
        from pptx import Presentation
        prs = Presentation(render(request('pptx', titleFontFamily='Georgia', headingFontFamily='Verdana', bodySize=20)))
        texts = [shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame]
        self.assertIn('Quarterly résumé', texts)
        self.assertIn('Overview', texts)
        self.assertTrue(any('Second line' in text for text in texts))
        table = next(shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table)
        self.assertEqual(table.cell(1, 1).text, '=SUM(A1:A2)')
        self.assertEqual(prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.name, 'Georgia')

    def test_pptx_keeps_wide_table_headers_without_data_rows(self):
        from pptx import Presentation
        prs = Presentation(render(request('pptx', [{'type': 'table', 'rows': [['Header%d' % n for n in range(9)]]}])))
        text = '\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn('Header8', text)

    def test_pptx_long_heading_requests_autofit(self):
        from pptx import Presentation
        from pptx.enum.text import MSO_AUTO_SIZE
        value = request('pptx', titleSize=40); value['document']['title'] = 'Long heading ' * 12
        prs = Presentation(render(value))
        self.assertEqual(prs.slides[0].shapes[0].text_frame.auto_size, MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE)

    def test_pdf_unicode_text_and_embedded_fonts(self):
        from pypdf import PdfReader
        pdf = PdfReader(render(request('pdf', fontFamily='Arial')))
        text = '\n'.join(page.extract_text() for page in pdf.pages)
        for expected in ('Quarterly résumé', 'Café', 'Ελληνικά', 'Привет', 'Second line', '=SUM(A1:A2)'):
            self.assertIn(expected, text)
        self.assertTrue(any('/FontFile2' in font.get_object().get('/FontDescriptor', {})
                            for page in pdf.pages for font in page['/Resources']['/Font'].values()
                            if '/FontDescriptor' in font.get_object()))

    def test_pdf_fails_unsupported_glyph_without_silent_replacement(self):
        value = request('pdf', [{'type': 'paragraph', 'text': '\U0001f9d1'}])
        result = renderer.handle_request(json.dumps(value).encode())
        self.assertFalse(result['ok'])
        self.assertIn('U+1F9D1', result['error'])

    def test_pdf_rejects_text_requiring_unavailable_shaping(self):
        for text in ('नमस्ते', 'مرحبا', 'שָׁלוֹם', 'q\u0307'):
            result = renderer.handle_request(json.dumps(request('pdf', [{'type': 'paragraph', 'text': text}])).encode())
            self.assertFalse(result['ok'], text)
            self.assertIn('shaping', result['error'])
            self.assertIn('DOCX', result['error'])

    def test_pdf_normalizes_ordinary_decomposed_accents_for_display(self):
        from pypdf import PdfReader
        pdf = PdfReader(render(request('pdf', [{'type': 'paragraph', 'text': 'Cafe\u0301'}])))
        self.assertIn('Café', pdf.pages[0].extract_text())

    def test_pdf_paginates_tall_cells_without_losing_text(self):
        from pypdf import PdfReader
        rows = [['Header'], ['\n'.join('Row %d' % n for n in range(150))]]
        pdf = PdfReader(render(request('pdf', [{'type': 'table', 'rows': rows}])))
        self.assertGreater(len(pdf.pages), 1)
        text = '\n'.join(page.extract_text() for page in pdf.pages)
        self.assertIn('Row 149', text)


if __name__ == '__main__':
    unittest.main()
