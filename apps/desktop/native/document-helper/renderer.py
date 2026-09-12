#!/usr/bin/env python3
"""Fixed trusted document operation. JSON accepts no executable input or paths."""
import base64
import io
import json
import math
from pathlib import Path
import re
import sys
import unicodedata

MAX_REQUEST = 1024 * 1024
MAX_SOURCE = 100 * 1024
MAX_OUTPUT = 4 * 1024 * 1024
MIME = {'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation', 'pdf': 'application/pdf'}
FONT_FILES = {'Arial': 'Arial.ttf', 'Times New Roman': 'Times New Roman.ttf', 'Georgia': 'Georgia.ttf',
              'Verdana': 'Verdana.ttf', 'Courier New': 'Courier New.ttf', 'Bradley Hand': 'Bradley Hand Bold.ttf',
              'Brush Script MT': 'Brush Script.ttf', 'Comic Sans MS': 'Comic Sans MS.ttf'}
SUBSETS = ('latin', 'latin-ext', 'cyrillic', 'cyrillic-ext', 'greek', 'greek-ext', 'vietnamese', 'devanagari')


class DocumentError(ValueError):
    pass


class OutputBuffer(io.BytesIO):
    def write(self, data):
        if self.tell() + len(data) > MAX_OUTPUT:
            raise DocumentError('Rendered artifact exceeds the 4 MB output limit. Split the content.')
        return super().write(data)


def keys(value, required, optional=()):
    if type(value) is not dict or not set(required) <= value.keys() or value.keys() - set(required) - set(optional):
        raise DocumentError('Invalid document request fields.')


def valid_text(text):
    if type(text) is not str or any((ord(c) < 32 and c not in '\n\t') or ord(c) == 127 or
                                   0xD800 <= ord(c) <= 0xDFFF or ord(c) in (0xFFFE, 0xFFFF) for c in text):
        raise DocumentError('Document text must contain valid Unicode without control characters.')
    return len(text.encode('utf-8'))


def validate(request):
    keys(request, ('version', 'operation', 'document'))
    if type(request['version']) is not int or request['version'] != 1 or request['operation'] != 'render_document':
        raise DocumentError('Unsupported document operation or protocol version.')
    document = request['document']
    keys(document, ('format', 'title', 'blocks'), ('typography',))
    if type(document['format']) is not str or document['format'] not in MIME:
        raise DocumentError('Unsupported document format.')
    valid_text(document['title'])
    if len(document['title']) > 160:
        raise DocumentError('Artifact title must be at most 160 characters.')
    typography = document.get('typography', {})
    keys(typography, (), ('titleSize', 'headingSize', 'bodySize', 'fontFamily', 'titleFontFamily', 'headingFontFamily', 'bodyFontFamily'))
    for key, value in typography.items():
        if key.endswith('Size'):
            if type(value) not in (float, int) or not math.isfinite(value) or not 8 <= value <= (24 if key == 'bodySize' else 40):
                raise DocumentError('Invalid document typography size.')
        elif type(value) is not str or value not in FONT_FILES:
            raise DocumentError('Choose a supported document font family.')
    blocks = document['blocks']
    if type(blocks) is not list or not blocks:
        raise DocumentError('Artifact content is required.')
    if len(blocks) > 4000:
        raise DocumentError('Artifact content exceeds the 4000 blocks limit.')
    source_size = 0; row_count = 0
    for block in blocks:
        if type(block) is not dict or type(block.get('type')) is not str:
            raise DocumentError('Invalid document block.')
        kind = block['type']
        if kind == 'table':
            keys(block, ('type', 'rows')); rows = block['rows']
            if type(rows) is not list or not rows or type(rows[0]) is not list or not rows[0]:
                raise DocumentError('Document tables require nonempty rows.')
            columns = len(rows[0])
            if columns > 30:
                raise DocumentError('Artifact tables support at most 30 columns.')
            row_count += len(rows)
            if row_count > 2000:
                raise DocumentError('Artifact tables exceed the 2000 rows limit.')
            for row in rows:
                if type(row) is not list or len(row) != columns:
                    raise DocumentError('Document table rows must have equal column counts.')
                source_size += sum(valid_text(cell) for cell in row)
        elif kind in ('heading', 'paragraph', 'code', 'list'):
            fields = ('type', 'text') + (('level',) if kind == 'heading' else ('ordered', 'marker') if kind == 'list' else ())
            keys(block, fields); source_size += valid_text(block['text'])
            if kind == 'heading' and (type(block['level']) is not int or not 1 <= block['level'] <= 6):
                raise DocumentError('Heading levels must be integers from 1 to 6.')
            if kind == 'list':
                if type(block['ordered']) is not bool or type(block['marker']) is not str or not (
                    re.fullmatch(r'[0-9]{1,100}\.', block['marker']) if block['ordered'] else block['marker'] == '•'):
                    raise DocumentError('Invalid document list marker.')
                source_size += valid_text(block['marker'])
        else:
            raise DocumentError('Unsupported document block type.')
        if source_size > MAX_SOURCE:
            raise DocumentError('Artifact content exceeds the 100 KB source limit.')
    return document


def family(typography, role, fallback='Aptos'):
    return typography.get(role + 'FontFamily', typography.get('fontFamily', fallback))


def block_text(block):
    if block['type'] == 'table':
        return '\n'.join('\t'.join(row) for row in block['rows'])
    return (block['marker'] + ' ' if block['type'] == 'list' else '') + block['text']


def render_docx(title, blocks, typography):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document(); document.core_properties.author = 'zQ'; document.core_properties.title = title
    section = document.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(2 / 3)
    roles = [('Normal', 'body', typography.get('bodySize', 11)), ('Title', 'title', typography.get('titleSize', 26))]
    roles += [('Heading ' + str(level), 'heading', max(8, typography.get('headingSize', 20) - (level - 1) * 2)) for level in range(1, 7)]
    role_sizes = {name: size for name, _, size in roles}

    def explicit_font(font, role, size, code=False):
        name = 'Courier New' if code else family(typography, role)
        font.name = name; font.size = Pt(size); font.color.rgb = RGBColor.from_string('202124')
        properties = font._element.get_or_add_rPr()
        fonts = properties.get_or_add_rFonts()
        for attribute in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + attribute), name)

    for name, role, size in roles:
        explicit_font(document.styles[name].font, role, size)
    # The bundled Word template includes theme references on linked character
    # styles and document defaults. Some Office previews prioritize these over
    # explicit families, so remove theme overrides throughout the fixed styles.
    for element in document.styles.element.iter():
        for attribute in list(element.attrib):
            if 'theme' in attribute.lower():
                del element.attrib[attribute]
    defaults = document.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    for tag, attributes in (
        ('rFonts', {key: family(typography, 'body') for key in ('ascii', 'hAnsi', 'eastAsia', 'cs')}),
        ('sz', {'val': str(round(role_sizes['Normal'] * 2))}),
        ('szCs', {'val': str(round(role_sizes['Normal'] * 2))}),
        ('color', {'val': '202124'}),
    ):
        element = defaults.find(qn('w:' + tag))
        if element is None:
            element = OxmlElement('w:' + tag); defaults.append(element)
        for attribute, value in attributes.items():
            element.set(qn('w:' + attribute), value)
    title_paragraph = document.add_paragraph(title, 'Title')
    for run in title_paragraph.runs:
        explicit_font(run.font, 'title', role_sizes['Title'])
    for block in blocks:
        kind = block['type']
        if kind == 'table':
            table = document.add_table(rows=0, cols=len(block['rows'][0])); table.style = 'Table Grid'
            for index, row in enumerate(block['rows']):
                cells = table.add_row().cells
                if index == 0:
                    table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for cell, text in zip(cells, row):
                    cell.text = text
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            explicit_font(run.font, 'body', role_sizes['Normal'])
                    if index == 0:
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'ECEEF0'); cell._tc.get_or_add_tcPr().append(shade)
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            document.add_paragraph()
        else:
            paragraph = document.add_paragraph(style='Heading ' + str(block['level']) if kind == 'heading' else 'Normal')
            run = paragraph.add_run(block_text(block))
            explicit_font(run.font, 'heading' if kind == 'heading' else 'body',
                          role_sizes['Heading ' + str(block['level'])] if kind == 'heading' else role_sizes['Normal'], kind == 'code')
            if kind == 'list':
                paragraph.paragraph_format.left_indent = Inches(.25); paragraph.paragraph_format.first_line_indent = Inches(-.125)
            paragraph.paragraph_format.space_after = Pt(7)
    output = OutputBuffer(); document.save(output)
    return output.getvalue()


def render_xlsx(title, blocks, typography):
    import mimetypes
    # openpyxl creates a MimeTypes database while importing its manifest. Use
    # Python's built-in mappings only; machine configuration is outside the
    # sandbox and must not influence a fixed document operation.
    mimetypes.knownfiles = []
    mimetypes.init(files=[])
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook(); workbook.properties.creator = 'zQ'; workbook.properties.title = title
    sheet = workbook.active; sheet.title = 'Content'; count = 0

    def add(values, heading=False, role='body', size=None):
        nonlocal count
        count += 1
        if count > 2000:
            raise DocumentError('Spreadsheet exceeds the 2000 rows limit. Split the content.')
        if len(values) > 30:
            raise DocumentError('Spreadsheet supports at most 30 columns.')
        for column, value in enumerate(values, 1):
            if len(value.encode('utf-16-le')) // 2 > 32767:
                raise DocumentError('Spreadsheet cells support at most 32767 characters.')
            cell = sheet.cell(count, column); cell.value = value
            cell.data_type = 's'  # Override formula AND error inference; source is always literal.
            cell.number_format = '@'
            cell.font = Font(name=family(typography, role), size=size or typography.get('bodySize', 11), bold=heading)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if heading:
                cell.fill = PatternFill('solid', fgColor='FFECEEF0')
    add([title], True, 'title', typography.get('titleSize', 11))
    for block in blocks:
        if block['type'] == 'table':
            for index, row in enumerate(block['rows']):
                add(row, index == 0, size=typography.get('headingSize', 11) if index == 0 else None)
        else:
            heading = block['type'] == 'heading'
            for line in block_text(block).split('\n'):
                add(line.split('\t'), heading, 'heading' if heading else 'body', typography.get('headingSize', 11) if heading else None)
    for column in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 28 if sheet.max_column > 1 else 95
    sheet.freeze_panes = 'A2'; sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    output = OutputBuffer(); workbook.save(output)
    return output.getvalue()


def wrap_characters(text, limit):
    lines = []
    for rest in text.split('\n'):
        while len(rest) > limit:
            boundary = rest[:limit].rfind(' '); count = boundary if boundary > limit // 2 else limit
            lines.append(rest[:count]); rest = rest[count:].lstrip()
        lines.append(rest)
    return lines


def render_pptx(title, blocks, typography):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE
    prs = Presentation(); prs.slide_width = Inches(13.333333); prs.slide_height = Inches(7.5)
    prs.core_properties.author = 'zQ'; prs.core_properties.title = title
    heading = title; title_slide = True; continuation = 0; section_used = False
    body_size = typography.get('bodySize', 18); budget = max(9, int(14 * 18 / body_size)); pending = []

    def style_paragraph(paragraph, role, size, bold=False):
        for run in paragraph.runs:
            run.font.name = family(typography, role); run.font.size = Pt(size)
            run.font.bold = bold; run.font.color.rgb = RGBColor.from_string('202124')

    def textbox(slide, text, x, y, w, h, role, size, bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame; frame.clear(); frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0; frame.text = text
        for paragraph in frame.paragraphs:
            style_paragraph(paragraph, role, size, bold)
        return shape

    def new_slide():
        nonlocal continuation, section_used
        if len(prs.slides) >= 80:
            raise DocumentError('Presentation exceeds the 80 slides limit. Split the content.')
        slide = prs.slides.add_slide(prs.slide_layouts[6]); section_used = True; role = 'title' if title_slide else 'heading'
        textbox(slide, heading + (' (continued)' if continuation else ''), .6, .4, 12.1, 1, role, typography.get(role + 'Size', 26), True)
        continuation += 1; textbox(slide, str(len(prs.slides)), 12, 7.05, .6, .2, 'body', 10)
        return slide

    def flush(all_lines=True):
        while len(pending) >= budget or (all_lines and pending):
            text = '\n'.join(pending[:budget]); del pending[:budget]
            textbox(new_slide(), text, .65, 1.6, 12, 5.15, 'body', body_size)

    for block in blocks:
        if block['type'] == 'heading':
            flush()
            if not section_used:
                new_slide()
            heading = block['text']
            if len(heading) > 160:
                raise DocumentError('Presentation section headings must be at most 160 characters.')
            title_slide = False; continuation = 0; section_used = False
        elif block['type'] == 'table' and len(block['rows'][0]) <= 6:
            flush(); columns = len(block['rows'][0]); limit = int(12 / columns * 8); rows = []
            for row in block['rows']:
                wrapped = [wrap_characters(cell, limit) for cell in row]
                for start in range(0, max(map(len, wrapped)), 3):
                    rows.append(['\n'.join(cell[start:start + 3]) for cell in wrapped])
            while rows:
                group = []; lines = 0
                while rows and lines + max(len(cell.split('\n')) for cell in rows[0]) + 1 <= 15:
                    row = rows.pop(0); lines += max(len(cell.split('\n')) for cell in row) + 1; group.append(row)
                table = new_slide().shapes.add_table(len(group), columns, Inches(.65), Inches(1.6), Inches(12), Inches(4.9)).table
                heights = [max(len(cell.split('\n')) for cell in row) + 1 for row in group]
                for index, row in enumerate(group):
                    table.rows[index].height = Inches(4.9 * heights[index] / sum(heights))
                    for column, text in enumerate(row):
                        cell = table.cell(index, column); cell.text = text
                        for paragraph in cell.text_frame.paragraphs:
                            style_paragraph(paragraph, 'body', 16)
        else:
            if block['type'] == 'table':
                text = '\n\n'.join('\n'.join(f'{block["rows"][0][i]}: {value}' for i, value in enumerate(row)) for row in block['rows'][1:])
                if len(block['rows']) == 1:
                    text = ' | '.join(block['rows'][0])
            else:
                text = block_text(block)
            pending.extend(wrap_characters(text, int(85 * 18 / body_size)) + ['']); flush(False)
    flush()
    if not section_used:
        new_slide()
    output = OutputBuffer(); prs.save(output)
    return output.getvalue()


def render_pdf(title, blocks, typography):
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    def display_text(text):
        # Compose ordinary accents for reliable glyph placement. Keep the host's
        # original source untouched; Office formats also retain original bytes.
        text = unicodedata.normalize('NFC', text)
        for character in text:
            code = ord(character)
            if (unicodedata.category(character).startswith('M') or
                unicodedata.bidirectional(character) in ('R', 'AL', 'AN', 'RLE', 'RLO', 'RLI', 'LRE', 'LRO', 'LRI', 'FSI', 'PDI', 'PDF') or
                0x0900 <= code <= 0x109F or 0x1100 <= code <= 0x11FF or
                0x1780 <= code <= 0x1CFF or 0xA800 <= code <= 0xABFF or code in (0x200C, 0x200D)):
                raise DocumentError('This PDF text requires complex-script shaping that the document helper does not support yet. Use DOCX for this text.')
        return text

    title = display_text(title)
    blocks = [{**block, 'rows': [[display_text(cell) for cell in row] for row in block['rows']]}
              if block['type'] == 'table' else {**block, 'text': display_text(block['text'])} for block in blocks]

    def register(name, path):
        try:
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(path)))
            return name, pdfmetrics.getFont(name)
        except Exception as error:
            raise DocumentError('A required PDF font is unavailable. Rebuild the document helper or choose another font.') from error

    resource_dir = Path(__file__).resolve().parent / 'fonts'
    fonts = [register('zq-noto-' + subset, resource_dir / ('noto-sans-' + subset + '-400-normal.ttf')) for subset in SUBSETS]
    selected = {}
    for role in ('title', 'heading', 'body'):
        name = family(typography, role, None)
        if name and name not in selected:
            selected[name] = register('zq-' + name.replace(' ', '-'), Path('/System/Library/Fonts/Supplemental') / FONT_FILES[name])
    glyphs = {}

    def glyph(character, role):
        name = family(typography, role, None); key = (character, name)
        if key not in glyphs:
            candidates = ([selected[name]] if name else []) + fonts
            found = next((font for font in candidates if font[1].face.charToGlyph.get(ord(character), 0) != 0), None)
            if not found:
                raise DocumentError(f'PDF font does not support “{character}” (U+{ord(character):X}). Use DOCX for this text.')
            glyphs[key] = found[0]
        return glyphs[key]

    for text, role in [(title, 'title')] + [(block_text(b), 'heading' if b['type'] == 'heading' else 'body') for b in blocks]:
        for character in text.replace('\n', ' ').replace('\t', ' '):
            glyph(character, role)
    output = OutputBuffer(); canvas = Canvas(output, pagesize=(612, 792), pageCompression=1)
    canvas.setTitle(title); canvas.setAuthor('zQ'); canvas.setCreator('zQ'); y = 744; pages = 1

    def ensure(height):
        nonlocal y, pages
        if y - height < 48:
            pages += 1
            if pages > 200:
                raise DocumentError('PDF exceeds the 200 pages limit. Split the content.')
            canvas.showPage(); y = 744

    def width(text, size, role):
        return sum(pdfmetrics.stringWidth(c, glyph(c, role), size) for c in text)

    def wrap(text, available, size, role):
        result = []
        for paragraph in text.replace('\t', '    ').split('\n'):
            line = ''
            for word in re.findall(r'\S+\s*|\s+', paragraph) or ['']:
                if line and width(line + word, size, role) > available:
                    result.append(line.rstrip()); line = ''
                if width(word, size, role) <= available:
                    line += word; continue
                clusters = []
                for character in word:
                    if clusters and unicodedata.combining(character):
                        clusters[-1] += character
                    else:
                        clusters.append(character)
                for cluster in clusters:
                    if line and width(line + cluster, size, role) > available:
                        result.append(line.rstrip()); line = ''
                    line += cluster
            result.append(line.rstrip())
        return result

    def draw(text, x, top, size, role):
        runs = []
        for c in text:
            font = glyph(c, role)
            if runs and runs[-1][0] == font:
                runs[-1][1] += c
            else:
                runs.append([font, c])
        canvas.setFillColorRGB(32 / 255, 33 / 255, 36 / 255)
        for font, text in runs:
            canvas.setFont(font, size); canvas.drawString(x, top - size, text)
            x += pdfmetrics.stringWidth(text, font, size)

    def paragraph(text, size=11, indent=0, role='body'):
        nonlocal y
        for line in wrap(text, 516 - indent, size, role):
            ensure(size * 1.5); draw(line, 48 + indent, y, size, role); y -= size * 1.5
        y -= 7

    paragraph(title, typography.get('titleSize', 23), role='title')
    for block in blocks:
        kind = block['type']
        if kind != 'table':
            if kind == 'heading':
                ensure(65); y -= 7
            size = (max(8, typography['headingSize'] - (block['level'] - 1) * 2) if 'headingSize' in typography else max(12, 20 - block['level'] * 2)) if kind == 'heading' else 9 if kind == 'code' else typography.get('bodySize', 11)
            paragraph(block_text(block), size, 10 if kind == 'list' else 0, 'heading' if kind == 'heading' else 'body')
        elif len(block['rows'][0]) > 8:
            paragraph(' | '.join(block['rows'][0]), 10)
            for row in block['rows'][1:]:
                paragraph('\n'.join(f'{block["rows"][0][i]}: {text}' for i, text in enumerate(row)), 10)
        else:
            cell_width = 516 / len(block['rows'][0])
            for index, row in enumerate(block['rows']):
                wrapped = [wrap(cell, cell_width - 12, 9, 'body') for cell in row]
                count = max(map(len, wrapped)); start = 0
                while start < count:
                    ensure(26); take = min(count - start, max(1, int((y - 48 - 10) / 14))); height = take * 14 + 10
                    for column, lines in enumerate(wrapped):
                        x = 48 + column * cell_width
                        canvas.setFillColorRGB(*((.925, .933, .941) if index == 0 else (1, 1, 1)))
                        canvas.setStrokeColorRGB(.788, .804, .82); canvas.rect(x, y - height, cell_width, height, fill=1, stroke=1)
                        for line in range(take):
                            draw(lines[start + line] if start + line < len(lines) else '', x + 6, y - 5 - line * 14, 9, 'body')
                    y -= height; start += take
            y -= 12
    canvas.save()
    return output.getvalue()


def unique_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise DocumentError('Duplicate request fields are not allowed.')
        result[key] = value
    return result


def invalid_constant(_):
    raise DocumentError('Invalid JSON number.')


def handle_request(raw):
    try:
        if len(raw) > MAX_REQUEST:
            raise DocumentError('Document request exceeds the 1 MB limit.')
        request = json.loads(raw, object_pairs_hook=unique_object, parse_constant=invalid_constant)
        document = validate(request); title = re.sub(r'\s+', ' ', document['title']).strip() or 'Untitled artifact'
        format = document['format']
        render = {'docx': render_docx, 'xlsx': render_xlsx, 'pptx': render_pptx, 'pdf': render_pdf}[format]
        data = render(title, document['blocks'], document.get('typography', {}))
        if not data or len(data) > MAX_OUTPUT:
            raise DocumentError('Rendered artifact exceeds the 4 MB output limit.')
        name = re.sub(r'[<>:"/\\|?*\u202a-\u202e\u2066-\u2069]', '-', title).strip('. \t\n')[:140] or 'artifact'
        while len(name.encode('utf-8')) > 240:
            name = name[:-1]
        name = name.rstrip('. ') or 'artifact'
        return {'version': 1, 'ok': True, 'file': {'name': name + '.' + format, 'mime': MIME[format],
                                                 'data': base64.b64encode(data).decode('ascii')}, 'warnings': []}
    except (DocumentError, UnicodeError, json.JSONDecodeError, RecursionError) as error:
        return {'version': 1, 'ok': False, 'error': str(error)[:500] or 'Invalid document request.'}
    except Exception:
        return {'version': 1, 'ok': False, 'error': 'Document rendering failed. Check the document helper installation and content.'}


def main():
    try:
        if len(sys.argv) == 2:
            # argv is selected by the native supervisor, never by JSON input.
            with open(sys.argv[1], 'rb') as source:
                raw = source.read(MAX_REQUEST + 1)
        elif len(sys.argv) == 1:
            raw = sys.stdin.buffer.read(MAX_REQUEST + 1)
        else:
            raise DocumentError('Invalid renderer invocation.')
        result = handle_request(raw)
    except Exception:
        result = {'version': 1, 'ok': False, 'error': 'Unable to read the document request.'}
    sys.stdout.write(json.dumps(result, ensure_ascii=True, separators=(',', ':')) + '\n')


if __name__ == '__main__':
    main()
