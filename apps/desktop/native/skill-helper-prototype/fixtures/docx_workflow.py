"""Synthetic DOCX round-trip fixture. Execute only through the isolated helper."""

import json
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Pt


output = Path("output")
output.mkdir(exist_ok=True)
destination = output / "document-edited.docx"

with tempfile.TemporaryDirectory(prefix="docx-fixture-", dir=".") as scratch:
    original_path = Path(scratch) / "original.docx"
    document = Document()
    document.add_heading("Synthetic document fixture", level=1)
    editable = document.add_paragraph()
    editable.add_run("Status: ").bold = True
    editable.add_run("Draft").italic = True
    preserved = document.add_paragraph("Preserve this unrelated paragraph exactly.")
    preserved.runs[0].font.size = Pt(11)
    document.sections[0].header.paragraphs[0].text = "Synthetic fixture header"
    document.sections[0].footer.paragraphs[0].text = "Synthetic fixture footer"
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Item", "Count"), ("Original", "7")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.save(original_path)

    edited = Document(original_path)
    edited.paragraphs[1].runs[1].text = "Reviewed"
    edited.save(destination)

    reopened = Document(destination)
    assert reopened.paragraphs[0].text == "Synthetic document fixture"
    assert reopened.paragraphs[1].text == "Status: Reviewed"
    assert reopened.paragraphs[1].runs[0].bold is True
    assert reopened.paragraphs[1].runs[1].italic is True
    assert reopened.paragraphs[2].text == "Preserve this unrelated paragraph exactly."
    assert reopened.paragraphs[2].runs[0].font.size == Pt(11)
    assert reopened.sections[0].header.paragraphs[0].text == "Synthetic fixture header"
    assert reopened.sections[0].footer.paragraphs[0].text == "Synthetic fixture footer"
    assert [[cell.text for cell in row.cells] for row in reopened.tables[0].rows] == [
        ["Item", "Count"], ["Original", "7"]
    ]

print(json.dumps({
    "fixture": "docx", "status": "passed", "output": str(destination),
    "checks": ["create", "reopen", "edit targeted run", "preserve unrelated paragraph",
               "preserve run formatting", "preserve header and footer", "preserve table"],
    "rendered": False,
}))
