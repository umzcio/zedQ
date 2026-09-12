"""Synthetic PPTX shape/slide round-trip with a generated Pillow image."""

import json
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches


output = Path("output")
output.mkdir(exist_ok=True)
destination = output / "presentation-edited.pptx"

with tempfile.TemporaryDirectory(prefix="pptx-fixture-", dir=".") as scratch:
    scratch = Path(scratch)
    image_path = scratch / "synthetic.png"
    image = Image.new("RGB", (160, 100), "white")
    ImageDraw.Draw(image).rectangle((20, 20, 140, 80), fill=(36, 78, 110))
    image.save(image_path)
    image.close()
    image_bytes = image_path.read_bytes()

    original_path = scratch / "original.pptx"
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    title = first.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8), Inches(0.8))
    title.name = "Editable title"
    title.text = "Draft presentation"
    text = first.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(7), Inches(0.8))
    text.name = "Unrelated text"
    text.text = "Preserve this shape exactly."
    rectangle = first.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(3), Inches(2), Inches(1))
    rectangle.name = "Unrelated rectangle"
    rectangle.fill.solid()
    rectangle.fill.fore_color.rgb = RGBColor(36, 78, 110)
    picture = first.shapes.add_picture(str(image_path), Inches(4), Inches(3), width=Inches(2))
    picture.name = "Unrelated picture"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = "Preserved second slide"
    deck.save(original_path)

    edited = Presentation(original_path)
    target = next(shape for shape in edited.slides[0].shapes if shape.name == "Editable title")
    target.text_frame.paragraphs[0].runs[0].text = "Reviewed presentation"
    third = edited.slides.add_slide(edited.slide_layouts[6])
    third.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = "Added third slide"
    edited.save(destination)

    reopened = Presentation(destination)
    assert len(reopened.slides) == 3
    shapes = {shape.name: shape for shape in reopened.slides[0].shapes}
    assert len(shapes) == 4
    assert shapes["Editable title"].text == "Reviewed presentation"
    assert shapes["Unrelated text"].text == "Preserve this shape exactly."
    assert shapes["Unrelated rectangle"].fill.fore_color.rgb == RGBColor(36, 78, 110)
    assert shapes["Unrelated rectangle"].left == Inches(0.6)
    assert shapes["Unrelated picture"].image.blob == image_bytes
    assert reopened.slides[1].shapes[0].text == "Preserved second slide"
    assert reopened.slides[2].shapes[0].text == "Added third slide"

print(json.dumps({
    "fixture": "pptx", "status": "passed", "output": str(destination),
    "checks": ["create", "reopen", "edit targeted shape", "add slide",
               "preserve unrelated shapes", "preserve picture bytes", "preserve second slide"],
    "slideCount": 3, "rendered": False,
}))
