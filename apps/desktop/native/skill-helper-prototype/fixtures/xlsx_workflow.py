"""Synthetic XLSX formula-preservation fixture; does not calculate formulas."""

import json
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


output = Path("output")
output.mkdir(exist_ok=True)
destination = output / "workbook-formulas.xlsx"
formulas = {"C2": "=A2*B2", "C3": "=A3*B3", "C4": "=SUM(C2:C3)"}

with tempfile.TemporaryDirectory(prefix="xlsx-fixture-", dir=".") as scratch:
    original_path = Path(scratch) / "original.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inputs"
    for column, width in {"A": 12, "B": 14, "C": 20, "D": 4, "E": 48}.items():
        sheet.column_dimensions[column].width = width
    sheet.append(["Quantity", "Unit price", "Total"])
    sheet.append([2, 7.5, formulas["C2"]])
    sheet.append([3, 4, formulas["C3"]])
    sheet["C4"] = formulas["C4"]
    sheet["E1"] = "Preserve this unrelated text exactly."
    sheet["E1"].font = Font(bold=True, color="005A9C")
    sheet["A2"].fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    sheet["B2"].number_format = "$0.00"
    workbook.create_sheet("Unrelated")["B3"] = "Preserved second sheet"
    workbook.save(original_path)
    workbook.close()

    edited = load_workbook(original_path)
    edited["Inputs"]["A2"] = 5
    edited.save(destination)
    edited.close()

    reopened = load_workbook(destination, data_only=False)
    assert reopened.sheetnames == ["Inputs", "Unrelated"]
    assert reopened["Inputs"]["A2"].value == 5
    for coordinate, formula in formulas.items():
        assert reopened["Inputs"][coordinate].value == formula
        assert reopened["Inputs"][coordinate].data_type == "f"
    assert reopened["Inputs"]["E1"].value == "Preserve this unrelated text exactly."
    assert reopened["Inputs"]["E1"].font.bold is True
    assert reopened["Inputs"]["A2"].fill.fgColor.rgb == "00FFF2CC"
    assert reopened["Inputs"]["B2"].number_format == "$0.00"
    assert reopened["Unrelated"]["B3"].value == "Preserved second sheet"
    reopened.close()

    values = load_workbook(destination, data_only=True)
    caches = {coordinate: values["Inputs"][coordinate].value for coordinate in formulas}
    assert all(value is None for value in caches.values()), caches
    values.close()

print(json.dumps({
    "fixture": "xlsx", "status": "passed", "output": str(destination),
    "checks": ["create", "reopen", "edit input", "preserve formulas",
               "preserve formatting", "preserve unrelated sheet"],
    "formulaCount": len(formulas), "calculatedCachesAvailable": False,
    "cachedValues": caches,
    "limitation": "openpyxl preserves formula strings but does not evaluate them. No calculation engine ran.",
    "rendered": False,
}))
